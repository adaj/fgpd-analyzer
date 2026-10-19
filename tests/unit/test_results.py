from __future__ import annotations

import json
from datetime import timedelta

import numpy as np
import pandas as pd
import pytest
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st
from typer.testing import CliRunner

from fgpd_analyzer import FGPDAnalyzer
from fgpd_analyzer.backends.base import LLMResponse
from fgpd_analyzer.backends.mock import MockBackend
from fgpd_analyzer.consensus import build_consensus, fleiss_kappa, iou, lower_median
from fgpd_analyzer.influence import UnitInfo, classify
from fgpd_analyzer.postprocess import Pattern, validate_patterns
from fgpd_analyzer.results import FGPDResults


def big_chat(n=24, group="g", gap_s=70):
    """Dyad plus an agent every 6th message; every message its own unit."""
    rows = []
    t0 = pd.Timestamp("2026-09-17 10:00:00")
    for i in range(n):
        who = "agente 1" if i % 6 == 0 else ("Ana" if i % 2 else "Bia")
        rows.append(
            {
                "Group": group,
                "Username": who,
                "Timestamp": str(t0 + pd.Timedelta(seconds=gap_s * i)),
                "Message": f"mensagem numero {i} sobre a tarefa",
            }
        )
    return pd.DataFrame(rows)


def analyzer(**kw):
    kw.setdefault("cache", "off")
    return FGPDAnalyzer(backend=kw.pop("backend_obj", None) or "mock", **kw)


@pytest.fixture
def pg():
    return FGPDAnalyzer().prepare(big_chat()).groups["g"]


def pat(run, s, e, ev, level=2, ante=None, goal="G1", up=None, inds=("share_idea",)):
    return Pattern("g", goal, run, s, e, tuple(ev), ante, tuple(inds), level, up, "j")


# ------------------------------------------------------------------ consensus


def test_iou_and_lower_median():
    assert iou((1, 4), (3, 6)) == pytest.approx(2 / 6)
    assert iou((1, 2), (3, 4)) == 0
    assert lower_median([2, 3, 4]) == 3 and lower_median([2, 3]) == 2 and lower_median([4]) == 4


def test_consensus_majority_and_unstable(pg):
    runs = {
        0: [pat(0, "u002", "u003", ["u002", "u003"], 2), pat(0, "u008", "u009", ["u008"], 4)],
        1: [pat(1, "u002", "u004", ["u002", "u003"], 3)],
        2: [pat(2, "u002", "u003", ["u003"], 4), pat(2, "u015", "u016", ["u015"], 1)],
    }
    res = build_consensus(runs, pg, "G1", 0.3, 0.5)
    assert len(res.consensus) == 1
    c = res.consensus[0]
    assert c.pattern.level == 3 and c.n_runs_detected == 3 and c.support == 1
    assert (c.pattern.start_id, c.pattern.end_id) == ("u002", "u003")
    assert c.pattern.evidence_ids == ("u002", "u003")  # each in >= half of members
    assert (c.level_min, c.level_max) == (2, 4)
    assert {u.pattern.start_id for u in res.unstable} == {"u008", "u015"}


def test_consensus_agent_uptake_needs_strict_majority(pg):
    runs = {
        r: [pat(r, "u002", "u003", ["u003"], 2, up=("u001" if r == 0 else None))] for r in range(2)
    }
    (c,) = build_consensus(runs, pg, "G1", 0.3, 0.5).consensus
    assert c.pattern.agent_uptake_id is None  # 1 of 2 is not > half
    runs[1] = [pat(1, "u002", "u003", ["u003"], 2, up="u001")]
    (c,) = build_consensus(runs, pg, "G1", 0.3, 0.5).consensus
    assert c.pattern.agent_uptake_id == "u001"


def test_consensus_single_run_and_empty(pg):
    res = build_consensus({0: [pat(0, "u002", "u003", ["u002"])]}, pg, "G1", 0.3, 0.5)
    assert len(res.consensus) == 1
    assert build_consensus({}, pg, "G1", 0.3, 0.5).consensus == []


def test_fleiss_kappa():
    perfect = np.array([[3, 0], [0, 3], [3, 0]], dtype=float)
    assert fleiss_kappa(perfect) == pytest.approx(1.0)
    assert fleiss_kappa(np.zeros((0, 2))) is None


# ------------------------------------------------------------------ influence


def _units():
    t0 = pd.Timestamp("2026-09-17 10:00:00")
    spec = [
        ("u001", True, 0),
        ("u002", False, 60),
        ("u003", False, 120),
        ("u004", True, 150),
        ("u005", False, 180),
        ("u006", True, 300),
        ("u007", False, 310),
    ]
    return {
        u: UnitInfo(
            u, i + 1, a, t0 + pd.Timedelta(seconds=s), t0 + pd.Timedelta(seconds=s), f"m{i + 1}"
        )
        for i, (u, a, s) in enumerate(spec)
    }


def test_influence_lag_boundaries():
    u = _units()
    lag = timedelta(seconds=60)  # u001 at 0 s, pattern starts u002 at 60 s: exactly at the lag
    inf = classify("u002", ("u003",), "u001", u, lag, "uptake")
    assert inf.agent_influence == "uptake" and inf.influenced and inf.agent_position == "before"
    assert inf.agent_lag_s == 120 and inf.agent_message_id == "m1"
    late = classify("u002", ("u003",), "u001", u, timedelta(seconds=59), "uptake")
    assert late.agent_influence == "none" and not late.influenced


def test_influence_within_temporal_and_after():
    u = _units()
    lag = timedelta(minutes=2)
    inside = classify("u002", ("u003", "u005"), "u004", u, lag, "uptake")
    assert inside.agent_influence == "uptake" and inside.agent_position == "within"
    assert inside.agent_lag_s == 30
    temporal = classify("u002", ("u003", "u005"), None, u, lag, "uptake")
    assert temporal.agent_influence == "temporal_only" and not temporal.influenced
    assert temporal.agent_unit_id == "u001"  # latest agent unit before the start
    assert classify("u002", ("u003", "u005"), None, u, lag, "temporal").influenced
    # cited agent after the last evidence is outside W
    after = classify("u002", ("u003",), "u006", u, lag, "uptake")
    assert after.agent_influence == "temporal_only"


# ------------------------------------------------------------------ end to end (mock)


def test_end_to_end_mock_outputs():
    res = analyzer().analyze(big_chat(36))
    f = res.frequencies
    assert set(f["goal"]) == {"G1", "G2", "G3", "G4"} and (f["status"] == "ok").all()
    p = res.patterns
    assert not p.empty and p["score"].isin([0.25, 0.5, 0.75, 1.0]).all()
    assert (p["start_min"] <= p["end_min"]).all()
    for _, row in f.iterrows():
        sub = p[(p.group == row.group) & (p.goal == row.goal)]
        assert row.frequency == pytest.approx(sub["score"].sum())
        assert row["count"] == len(sub)
    assert "text" not in res.messages.columns and "text" not in res.units.columns
    meta = res.run_metadata
    assert meta["n_calls"] == 12 and meta["backend"] == "mock"
    assert set(meta["stability"]) == {"G1", "G2", "G3", "G4"}
    assert meta["prompt_set_version"] == "v1" and len(meta["prompts"]) == 8


def test_no_overlap_within_goal():
    p = analyzer().analyze(big_chat(36)).patterns
    order = {f"u{i:03d}": i for i in range(1, 100)}
    for _, sub in p.groupby(["group", "goal"]):
        spans = sorted((order[r.start_unit_id], order[r.end_unit_id]) for r in sub.itertuples())
        assert all(a[1] < b[0] for a, b in zip(spans, spans[1:], strict=False))


@settings(max_examples=15, deadline=None, suppress_health_check=[HealthCheck.too_slow])
@given(every=st.integers(min_value=1, max_value=40), n=st.integers(min_value=14, max_value=40))
def test_timeseries_properties(every, n):
    res = analyzer().analyze(big_chat(n))
    ts = res.timeseries(f"{every}min")
    f = res.frequencies.set_index(["group", "goal"])
    for (g, goal), sub in ts.groupby(["group", "goal"]):
        vals = sub.sort_values("bin_index")
        assert vals["frequency_cum"].is_monotonic_increasing
        assert vals["count_cum"].is_monotonic_increasing
        assert vals["frequency_cum"].iloc[-1] == pytest.approx(f.loc[(g, goal), "frequency"])
        assert vals["student_messages_cum"].iloc[-1] == f.loc[(g, goal), "n_student_messages"]
    inc = res.timeseries(f"{every}min", cumulative=False)
    tot = inc.groupby(["group", "goal"])["frequency"].sum()
    for k, v in tot.items():
        assert v == pytest.approx(f.loc[k, "frequency"])


def test_duplicate_group_gives_identical_metrics():
    df = pd.concat([big_chat(30, "a"), big_chat(30, "b")], ignore_index=True)
    f = analyzer().analyze(df).frequencies
    cols = ["frequency", "count", "rate_per_100", "participation"]
    a = f[f.group == "a"][cols].reset_index(drop=True)
    b = f[f.group == "b"][cols].reset_index(drop=True)
    pd.testing.assert_frame_equal(a, b)


def test_row_shuffle_invariance():
    df = big_chat(30)
    a = analyzer().analyze(df).frequencies
    b = analyzer().analyze(df.sample(frac=1, random_state=3)).frequencies
    pd.testing.assert_frame_equal(a, b)


class Exploding(MockBackend):
    async def complete(self, request):  # noqa: D401
        raise AssertionError("network call made")


def test_dry_run_and_with_influence_make_no_calls():
    fa = FGPDAnalyzer(backend=Exploding(), cache="off")
    out = fa.analyze(big_chat(30), dry_run=True)
    assert out["groups"] == ["g"] and out["estimate"].n_calls == 12
    res = analyzer().analyze(big_chat(30))
    res2 = res.with_influence("10s")
    assert (
        res2.patterns["influenced"].sum()
        <= res.with_influence("10min").patterns["influenced"].sum()
    )
    assert res2.run_metadata["influence"]["lag_s"] == 10


def test_statuses_and_na_metrics():
    df = pd.concat([big_chat(30, "ok"), big_chat(5, "tiny")], ignore_index=True)

    def script(req):
        if req.goal == "G3":
            return LLMResponse("", "refusal")
        if req.goal == "G4" and req.run_idx == 0:
            return LLMResponse("{bad", "end_turn")
        return None

    res = FGPDAnalyzer(backend=MockBackend(script=script), cache="off").analyze(df)
    f = res.frequencies.set_index(["group", "goal"])
    assert f.loc[("ok", "G1"), "status"] == "ok"
    assert f.loc[("ok", "G3"), "status"] == "refused" and pd.isna(f.loc[("ok", "G3"), "frequency"])
    assert f.loc[("ok", "G4"), "status"] == "partial" and f.loc[("ok", "G4"), "n_runs_ok"] == 2
    assert f.loc[("tiny", "G1"), "status"] == "insufficient_data"
    assert pd.isna(f.loc[("tiny", "G1"), "frequency"])
    assert not res.timeseries().query("group == 'tiny'").shape[0]


def test_cache_modes(tmp_path):
    df = big_chat(30)
    mb = MockBackend()
    FGPDAnalyzer(backend=mb, cache="use", cache_dir=tmp_path).analyze(df)
    assert len(mb.calls) == 12
    mb2 = MockBackend()
    r = FGPDAnalyzer(backend=mb2, cache="use", cache_dir=tmp_path).analyze(df)
    assert len(mb2.calls) == 0 and r.run_metadata["n_cached"] == 12
    mb3 = MockBackend()
    FGPDAnalyzer(backend=mb3, cache="refresh", cache_dir=tmp_path).analyze(df)
    assert len(mb3.calls) == 12
    mb4 = MockBackend()
    FGPDAnalyzer(backend=mb4, cache="off", cache_dir=tmp_path).analyze(df)
    assert len(mb4.calls) == 12


def test_mock_and_anthropic_cache_namespaces_differ(tmp_path):
    from fgpd_analyzer.backends.mock import ReplayBackend

    df = big_chat(30)
    FGPDAnalyzer(backend=MockBackend(), cache="use", cache_dir=tmp_path).analyze(df)
    # a replay of "anthropic" responses must not be served the cached mock responses
    r = FGPDAnalyzer(backend=ReplayBackend([]), cache="use", cache_dir=tmp_path).analyze(df)
    assert (r.calls["status"] == "failed").all()


def test_save_load_roundtrip(tmp_path):
    res = analyzer().analyze(big_chat(30))
    for fmt in ("csv", "parquet"):
        d = tmp_path / fmt
        res.save(d, format=fmt, timeseries_every=("5min", "10min"))
        back = FGPDResults.load(d)
        pd.testing.assert_frame_equal(back.frequencies, res.frequencies, check_dtype=False)
        pd.testing.assert_frame_equal(
            back.timeseries("3min"), res.timeseries("3min"), check_dtype=False
        )
        meta = json.loads((d / "run_metadata.json").read_text())
        assert meta["influence"]["rule"] == "uptake"


def test_privacy_sentinel_in_outputs(tmp_path):
    df = big_chat(30)
    df.loc[df.Username == "Ana", "Username"] = "Sentinelius"
    res = analyzer(extra_names=["Zebedeu"]).analyze(df)
    d = tmp_path / "out"
    res.save(d)
    for f in d.iterdir():
        text = f.read_text(encoding="utf-8", errors="ignore")
        assert "Sentinelius" not in text and "Zebedeu" not in text, f.name


# ------------------------------------------------------------------ CLI


def test_cli_exit_codes(tmp_path):
    from fgpd_analyzer.cli import app

    path = tmp_path / "chats.csv"
    big_chat(30).to_csv(path, index=False)
    runner = CliRunner()
    assert runner.invoke(app, ["validate", str(path)]).exit_code == 0
    bad = tmp_path / "bad.csv"
    big_chat(30).drop(columns="Message").to_csv(bad, index=False)
    assert runner.invoke(app, ["validate", str(bad)]).exit_code == 2
    out = tmp_path / "res"
    r = runner.invoke(
        app, ["analyze", str(path), "-o", str(out), "--backend", "mock"], env={"FGPD_UNUSED": "1"}
    )
    assert r.exit_code == 0, r.output
    assert (out / "frequencies.csv").exists()
    assert runner.invoke(app, ["influence", str(out), "--lag", "30s"]).exit_code == 0
    assert runner.invoke(app, ["timeseries", str(out), "--every", "10min"]).exit_code == 0
    assert (out / "timeseries_10min.csv").exists()
    assert runner.invoke(app, ["estimate", str(path)]).exit_code == 0
    assert runner.invoke(app, ["prompts", "list"]).exit_code == 0
    assert (
        runner.invoke(app, ["analyze", str(bad), "-o", str(out), "--backend", "mock"]).exit_code
        == 2
    )


def test_cli_coding_sheet_check(tmp_path):
    from openpyxl import load_workbook

    from fgpd_analyzer.cli import app

    path = tmp_path / "chats.csv"
    big_chat(30).to_csv(path, index=False)
    sheet = tmp_path / "s.xlsx"
    runner = CliRunner()
    assert runner.invoke(app, ["coding-sheet", "export", str(path), str(sheet)]).exit_code == 0
    wb = load_workbook(sheet)
    wb["G2"].append(["u003", "u004", "u004", None, "rephrase", "3", None, "x"])  # no antecedent
    wb.save(sheet)
    r = runner.invoke(app, ["coding-sheet", "check", str(path), str(sheet)])
    assert r.exit_code == 4 and "drop_missing_antecedent" in r.output


# ------------------------------------------------------------------ M3 audit gaps


def _v(pg, goal, *items, apply=True):
    return validate_patterns(list(items), pg, goal, 0, apply=apply)


def R(**kw):
    base = {
        "start_id": "u002",
        "end_id": "u003",
        "evidence_ids": ["u003"],
        "antecedent_id": "u002",
        "indicators": ["rephrase"],
        "justification": "j",
        "level": 2,
        "agent_uptake_id": None,
    }
    base.update(kw)
    return base


def test_uppercase_ids_and_float_level(pg):
    (p,), _ = _v(
        pg,
        "G2",
        R(start_id="U002", end_id="U003", evidence_ids=["U003"], antecedent_id="U002", level=3.0),
    )
    assert (p.start_id, p.level) == ("u002", 3)


def test_more_flags(pg):
    (p,), qc = _v(pg, "G2", R(evidence_ids=["u001", "u003"], start_id="u001"))
    assert "evidence_not_student_removed" in p.qc_flags
    (p,), _ = _v(
        pg,
        "G1",
        R(
            antecedent_id="u005",
            start_id="u002",
            end_id="u003",
            evidence_ids=["u003"],
            indicators=["share_idea"],
        ),
    )
    assert p.antecedent_id is None and "antecedent_after_evidence_nulled" in p.qc_flags
    pats, qc = _v(pg, "G2", R(evidence_ids=["u002"]))
    assert pats == [] and qc.counts["drop_no_evidence"] == 1
    pats, qc = _v(pg, "G1", R(level=7, indicators=["share_idea"]))
    assert pats == [] and qc.counts["drop_invalid_level"] == 1
    # G1 with a same-speaker antecedent: nulled (codebook §3)
    (p,), _ = _v(
        pg,
        "G1",
        R(
            start_id="u003",
            end_id="u005",
            evidence_ids=["u005"],
            antecedent_id="u003",
            indicators=["share_idea"],
        ),
    )
    assert p.antecedent_id is None and "antecedent_same_speaker_nulled" in p.qc_flags


def test_merge_rechecks_rules_and_keeps_evidence(pg):
    # b's antecedent u002 is a's evidence: evidence must not shrink
    pats, _ = _v(
        pg,
        "G1",
        R(
            start_id="u002",
            end_id="u004",
            evidence_ids=["u002", "u004"],
            antecedent_id=None,
            indicators=["share_idea"],
        ),
        R(
            start_id="u003",
            end_id="u005",
            evidence_ids=["u005"],
            antecedent_id="u002",
            indicators=["share_idea"],
        ),
    )
    (m,) = pats
    assert set(m.evidence_ids) == {"u002", "u004", "u005"} and m.antecedent_id is None
    # chain of three overlapping patterns becomes one
    pats, qc = _v(
        pg,
        "G1",
        R(
            start_id="u002",
            end_id="u004",
            evidence_ids=["u002"],
            antecedent_id=None,
            indicators=["share_idea"],
        ),
        R(
            start_id="u004",
            end_id="u008",
            evidence_ids=["u005"],
            antecedent_id=None,
            indicators=["share_idea"],
        ),
        R(
            start_id="u008",
            end_id="u010",
            evidence_ids=["u009"],
            antecedent_id=None,
            indicators=["share_idea"],
            level=4,
        ),
    )
    # c alone is demoted 4 -> 2 (one unit, one speaker) before the merge
    assert len(pats) == 1 and pats[0].level == 2 and qc.counts["overlap_merged"] == 2


def test_report_only_finds_nested_overlaps(pg):
    pats, qc = _v(
        pg,
        "G1",
        R(
            start_id="u002",
            end_id="u009",
            evidence_ids=["u003"],
            antecedent_id=None,
            indicators=["share_idea"],
        ),
        R(
            start_id="u003",
            end_id="u004",
            evidence_ids=["u003"],
            antecedent_id=None,
            indicators=["share_idea"],
        ),
        R(
            start_id="u007",
            end_id="u008",
            evidence_ids=["u007"],
            antecedent_id=None,
            indicators=["share_idea"],
        ),
        apply=False,
    )
    assert qc.counts["overlap"] == 2
    pats, _ = _v(pg, "G1", R(start_id="u999", indicators=["share_idea"]), apply=False)
    assert pats[0].start_id == "u999" and "drop_unknown_span" in pats[0].qc_flags


def test_malformed_http_body_fails_one_call_only(pg):
    import anthropic
    import httpx2

    from fgpd_analyzer.backends.anthropic import AnthropicBackend

    def handler(request):
        return httpx2.Response(200, text="<html>proxy error</html>")

    client = anthropic.AsyncAnthropic(
        api_key="sk-test",
        max_retries=0,
        http_client=httpx2.AsyncClient(transport=httpx2.MockTransport(handler)),
    )
    res = FGPDAnalyzer(backend=AnthropicBackend(client=client), cache="off").analyze(big_chat(30))
    assert (res.calls["status"] == "failed").all() and len(res.calls) == 12


# ------------------------------------------------------------------ M4 audit fixes


def test_csv_roundtrip_keeps_string_ids(tmp_path):
    df = pd.concat([big_chat(30, "007"), big_chat(30, "NA")], ignore_index=True)
    df["message_id"] = [f"{i:03d}" for i in range(len(df))]
    res = analyzer().analyze(df)
    res.save(tmp_path)
    back = FGPDResults.load(tmp_path)
    assert set(back.frequencies["group"]) == {"007", "NA"}
    assert back.messages["message_id"].iloc[0] == "000"
    pd.testing.assert_frame_equal(back.frequencies, res.frequencies, check_dtype=False)


def test_not_selected_groups():
    df = pd.concat([big_chat(30, "a"), big_chat(30, "b")], ignore_index=True)
    f = analyzer().analyze(df, groups=["a"]).frequencies
    assert set(f[f.group == "b"]["status"]) == {"not_selected"}
    assert set(f[f.group == "a"]["status"]) == {"ok"}


def test_schema_fallback_is_replayable(tmp_path):
    from fgpd_analyzer.backends.mock import ReplayBackend
    from fgpd_analyzer.exceptions import BackendError

    class RejectsEnums(MockBackend):
        async def complete(self, request):
            schema = request.output_config["format"]["schema"]
            if "enum" in schema["$defs"]["unit_id"]:
                self.calls.append(request)
                raise BackendError("schema too complex", status=400)
            return await super().complete(request)

    df = big_chat(30)
    live = FGPDAnalyzer(backend=RejectsEnums(), cache="use", cache_dir=tmp_path).analyze(df)
    assert (live.calls["status"] == "ok").all()
    from fgpd_analyzer.cache import ResponseCache

    jsonl = ResponseCache(tmp_path).export_jsonl(tmp_path / "rec.jsonl")
    rep = FGPDAnalyzer(backend=ReplayBackend(jsonl, namespace="mock"), cache="off").analyze(df)
    assert (rep.calls["status"] == "ok").all()
    pd.testing.assert_frame_equal(rep.frequencies, live.frequencies)
    again = RejectsEnums()
    FGPDAnalyzer(backend=again, cache="use", cache_dir=tmp_path).analyze(df)
    assert again.calls == []  # straight to the cached fallback, no failing 400 call


def test_even_runs_keep_supported_evidence(pg):
    runs = {
        r: [pat(r, s, e, ["u004"] if r < 2 else ["u004", "u005"], 2)]
        for r, (s, e) in enumerate(
            [("u002", "u003"), ("u002", "u003"), ("u004", "u005"), ("u004", "u005")]
        )
    }
    for r in (0, 1):  # evidence must be inside the member span
        runs[r] = [pat(r, "u002", "u004", ["u004"], 2)]
    res = build_consensus(runs, pg, "G1", 0.3, 0.5)
    assert len(res.consensus) == 1
    assert "u004" in res.consensus[0].pattern.evidence_ids


def test_no_dispersion_flag_and_kappa_undefined(pg):
    res = build_consensus({0: [pat(0, "u002", "u003", ["u002"])]}, pg, "G1", 0.3, 0.5)
    assert "no_dispersion" in res.consensus[0].pattern.qc_flags
    assert fleiss_kappa(np.array([[3, 0], [3, 0]], dtype=float)) is None


def test_optional_antecedent_needs_majority(pg):
    runs = {
        0: [pat(0, "u004", "u006", ["u005"], 2, ante="u004")],
        1: [pat(1, "u004", "u006", ["u005"], 2)],
        2: [pat(2, "u004", "u006", ["u005"], 2)],
    }
    (c,) = build_consensus(runs, pg, "G1", 0.3, 0.5).consensus
    assert c.pattern.antecedent_id is None


def test_pseudonym_map_only_on_request(tmp_path):
    res = analyzer().analyze(big_chat(30))
    res.save(tmp_path / "a")
    assert not (tmp_path / "a" / "pseudonym_map.json").exists()
    from fgpd_analyzer.exceptions import PrivacyWarning

    with pytest.warns(PrivacyWarning):
        res.save(tmp_path / "b", include_pseudonym_map=True)
    m = json.loads((tmp_path / "b" / "pseudonym_map.json").read_text())
    assert m["g"]["Ana"].startswith("S")
    assert res.run_metadata["groups"]["g"]["status"]["G1"] == "ok"


def test_cli_influence_keeps_format_and_bad_rule(tmp_path):
    from fgpd_analyzer.cli import app

    path = tmp_path / "chats.csv"
    big_chat(30).to_csv(path, index=False)
    out = tmp_path / "res"
    runner = CliRunner()
    r = runner.invoke(
        app,
        [
            "analyze",
            str(path),
            "-o",
            str(out),
            "--backend",
            "mock",
            "--format",
            "parquet",
            "--every",
            "5min",
            "--every",
            "10min",
            "--cache",
            "off",
        ],
    )
    assert r.exit_code == 0, r.output
    assert runner.invoke(app, ["influence", str(out), "--lag", "30s"]).exit_code == 0
    assert not list(out.glob("*.csv"))
    assert {p.name for p in out.glob("timeseries_*")} == {
        "timeseries_5min.parquet",
        "timeseries_10min.parquet",
    }
    assert runner.invoke(app, ["influence", str(out), "--rule", "nope"]).exit_code == 2
    assert (
        runner.invoke(
            app, ["analyze", str(path), "-o", str(out), "--backend", "mock", "--effort", "huge"]
        ).exit_code
        == 2
    )


def test_coding_sheet_rows_are_sheet_rows(tmp_path):
    from openpyxl import load_workbook

    from fgpd_analyzer.coding import check_coding_sheet

    fa = FGPDAnalyzer()
    df = big_chat(30)
    sheet = fa.export_coding_sheet(df, tmp_path / "s.xlsx")
    wb = load_workbook(sheet)
    ws = wb["G2"]
    ws.append(["u003", "u004", "u004", "u003", "rephrase", "2", None, "ok"])
    ws.append([None] * 8)
    ws.append(["u005", "u006", "u006", None, "rephrase", "2", None, "no antecedent"])
    wb.save(sheet)
    rep = check_coding_sheet(fa.prepare(df), sheet)
    assert rep["row"].tolist() == [6] and "drop_missing_antecedent" in rep["problems"].iloc[0]
