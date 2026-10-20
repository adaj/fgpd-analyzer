from __future__ import annotations

import re
from pathlib import Path

import pytest

from fgpd_analyzer import FGPDAnalyzer, PromptError, PromptSchemaMismatch
from fgpd_analyzer.codebook import GOAL_INDICATORS, render_codebook_block
from fgpd_analyzer.coding import FORM_COLUMNS, read_coding_sheet
from fgpd_analyzer.config import GOALS
from fgpd_analyzer.prompts.registry import PROMPT_NAMES, PromptRegistry

REPO = Path(__file__).resolve().parents[2]


def test_builtin_prompts_and_variables():
    fa = FGPDAnalyzer()
    table = fa.list_prompts()
    assert table["name"].tolist() == list(PROMPT_NAMES)
    assert set(table["source"]) == {"builtin"}
    v = dict(zip(table["name"], table["variables"], strict=True))
    assert v["codebook_general"] == "justification_language, pattern_gap_min, pattern_gap_units"
    assert v["task_detect"] == "goal, goal_label"


def test_indicator_tables_match_constants():
    reg = PromptRegistry()
    for g in GOALS:
        text = reg.get(f"goal_{g}").text
        table = text.split("## Indicators", 1)[1].split("## Levels", 1)[0]
        codes = tuple(re.findall(r"^\| `([a-z_]+)` \|", table, flags=re.MULTILINE))
        assert codes == GOAL_INDICATORS[g], g


def test_examples_only_use_declared_indicators():
    reg = PromptRegistry()
    for g in GOALS:
        examples = reg.get(f"goal_{g}").text.split("## Examples", 1)[1]
        used = set(re.findall(r"`([a-z_]+)`", examples))
        assert used <= set(GOAL_INDICATORS[g]), (g, used - set(GOAL_INDICATORS[g]))


def test_rendering_uses_config():
    fa = FGPDAnalyzer(
        pattern_gap_units=4, pattern_gap="3min", justification_language="pt", chat_language="pt-BR"
    )
    general = fa.show_prompt("codebook_general", rendered=True)
    assert "4 units of any speaker, or more than 3 minutes" in general
    assert "written in Portuguese" in general
    assert "The chat is in Brazilian Portuguese" in fa.show_prompt(
        "chat_conventions", rendered=True
    )
    assert "{{" in fa.show_prompt("codebook_general")
    with pytest.raises(PromptError):
        fa.show_prompt("task_detect", rendered=True)
    assert "pattern of G3 (deepen)" in fa.show_prompt("task_detect", rendered=True, goal="G3")


def test_set_prompt_checks_variables_and_marks_custom():
    fa = FGPDAnalyzer()
    base_hash = fa.prompt_set_hash
    with pytest.raises(PromptSchemaMismatch):
        fa.set_prompt("task_detect", "Find all {{ goal }} patterns.")  # drops goal_label
    with pytest.raises(PromptSchemaMismatch):
        fa.set_prompt("system", "Hi {{ name }}")
    with pytest.raises(PromptError):
        fa.set_prompt("system", "{% if %}")
    with pytest.raises(PromptError):
        fa.set_prompt("nope", "x")
    p = fa.set_prompt("system", "You are a strict analyst.")
    assert p.source == "custom"
    assert fa.prompt_set_version.startswith("custom-")
    assert fa.prompt_set_hash != base_hash
    fa.reset_prompts(["system"])
    assert fa.prompt_set_version == "v1" and fa.prompt_set_hash == base_hash


def test_prompt_hash_stable_and_config_sensitive():
    assert FGPDAnalyzer().prompt_set_hash == FGPDAnalyzer().prompt_set_hash
    assert FGPDAnalyzer().prompt_set_hash != FGPDAnalyzer(unit_gap="30s").prompt_set_hash
    assert FGPDAnalyzer().prompt_set_hash == FGPDAnalyzer(influence_lag="5min").prompt_set_hash


def test_export_and_load_prompts_dir(tmp_path):
    fa = FGPDAnalyzer()
    files = fa.export_prompts(tmp_path)
    assert len(files) == len(PROMPT_NAMES)
    (tmp_path / "system.md.j2").write_text("Custom system prompt.\n", encoding="utf-8")
    fb = FGPDAnalyzer(prompts_dir=tmp_path)
    assert fb.get_prompt("system").source == "custom"
    assert fb.get_prompt("goal_G1").source == "builtin"  # identical text stays builtin
    fc = FGPDAnalyzer(prompts={"system": "Another one."})
    assert fc.get_prompt("system").text == "Another one."


def test_codebook_sections_are_the_rendered_prompts(tmp_path):
    fa = FGPDAnalyzer(justification_language="pt")
    text = fa.export_codebook(tmp_path / "cb.md").read_text(encoding="utf-8")
    assert fa.prompt_set_hash in text
    for name in PROMPT_NAMES:
        if name == "task_detect":
            for g in GOALS:
                assert fa.show_prompt(name, rendered=True, goal=g).strip() in text
        else:
            assert fa.show_prompt(name, rendered=True).strip("\n") in text, name
    block = render_codebook_block(fa.prompts, fa.config)
    for g in GOALS:
        assert fa.show_prompt(f"goal_{g}", rendered=True).strip("\n") in block


def test_committed_codebook_is_current(tmp_path):
    """CI guard (DESIGN K R1): docs/codebook.md must be regenerated after any prompt change."""
    committed = REPO / "docs" / "codebook.md"
    fresh = FGPDAnalyzer().export_codebook(tmp_path / "codebook.md").read_text(encoding="utf-8")
    assert committed.exists(), "run: python -m fgpd_analyzer.tools.regenerate_docs"
    assert committed.read_text(encoding="utf-8") == fresh, (
        "docs/codebook.md is stale; run: python -m fgpd_analyzer.tools.regenerate_docs"
    )


def test_coding_sheet_roundtrip(tmp_path, small_chat):
    from openpyxl import load_workbook

    fa = FGPDAnalyzer()
    path = fa.export_coding_sheet(small_chat, tmp_path / "sheet.xlsx", coder="coder-A")
    wb = load_workbook(path)
    assert wb.sheetnames == ["README", "transcript", "messages", "G1", "G2", "G3", "G4"]
    tr = wb["transcript"]
    header = [c.value for c in tr[1]]
    assert header == ["group", "unit_id", "time", "minutes", "speaker", "text"]
    texts = " ".join(str(c.value) for c in tr["F"])
    assert "Mariana" not in texts and "Joao" not in texts
    ws = wb["G2"]
    assert [c.value for c in ws[3]] == list(FORM_COLUMNS)
    ws.append(
        ["u003", "u004", "u004", "u002", "rephrase, content_reference", "3", None, "S2 retoma u002"]
    )
    wb.save(path)
    coded = read_coding_sheet(path)
    assert len(coded) == 1
    row = coded.iloc[0]
    assert row["goal"] == "G2" and row["group"] == "g1" and row["level"] == 3
    assert row["evidence_ids"] == ["u004"] and row["antecedent_id"] == "u002"
    assert row["indicators"] == ["rephrase", "content_reference"]
    assert row["agent_uptake_id"] is None and row["coder"] == "coder-A"


def test_prompt_set_v1_is_frozen():
    """v1 was frozen for the 0.1.0 release. Any change to a built-in prompt must go to a
    new prompt set (prompts/v2) with a CHANGELOG entry, never edit v1 in place."""
    import json
    from importlib import resources

    manifest = json.loads(
        (resources.files("fgpd_analyzer.prompts") / "v1" / "MANIFEST.json").read_text()
    )
    reg = PromptRegistry()
    changed = [n for n, sha in manifest["prompts"].items() if reg.get(n).sha256 != sha]
    assert not changed, f"frozen prompt set v1 was edited: {changed}; create prompts/v2"
